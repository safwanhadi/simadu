from datetime import date
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from myaccount.models import ProfilSDM, Users
from strukturorg.models import InstansiDaerah, SatuanKerjaInduk, UnitOrganisasi

from .models import RiwayatCuti, RiwayatPenempatan


class DataCutiPegawaiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = Users.objects.create_superuser(
            email='admin-data-cuti@example.com',
            password='Password-123!',
        )
        cls.pegawai = Users.objects.create_user(
            email='pegawai-data-cuti@example.com',
            first_name='Siti',
            last_name='Sehat',
            password='Password-123!',
        )
        ProfilSDM.objects.create(
            user=cls.pegawai,
            nip='198001012010012001',
            no_hp='08123456789',
            email_pribadi='siti@example.com',
        )
        instansi = InstansiDaerah.objects.create(instansi='Pemerintah Daerah')
        satker = SatuanKerjaInduk.objects.create(
            instansi_daerah=instansi,
            satuan_kerja='Rumah Sakit',
        )
        unor = UnitOrganisasi.objects.create(
            satker_induk=satker,
            unor='Bidang Pelayanan',
        )
        RiwayatPenempatan.objects.create(
            pegawai=cls.pegawai,
            penempatan_level1=unor,
            status=True,
        )
        cls.cuti_tahunan = RiwayatCuti.objects.create(
            pegawai=cls.pegawai,
            jenis_cuti='Cuti Tahunan',
            tgl_mulai_cuti=date(2026, 1, 5),
            tgl_akhir_cuti=date(2026, 1, 7),
            lama_cuti=3,
            tahun_cuti=2026,
        )
        cls.cuti_sakit = RiwayatCuti.objects.create(
            pegawai=cls.pegawai,
            jenis_cuti='Cuti Sakit',
            tgl_mulai_cuti=date(2025, 2, 1),
            tgl_akhir_cuti=date(2025, 2, 2),
            lama_cuti=2,
            tahun_cuti=2025,
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_daftar_memuat_semua_jenis_cuti_dan_penempatan(self):
        response = self.client.get(reverse('riwayat_urls:riwayat_cuti_melahirkan'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context['cuti_list']), [
            self.cuti_tahunan,
            self.cuti_sakit,
        ])
        self.assertContains(response, 'Bidang Pelayanan')
        self.assertContains(response, 'Cuti Tahunan')
        self.assertContains(response, 'Cuti Sakit')

    def test_filter_jenis_tahun_dan_pegawai_diterapkan(self):
        response = self.client.get(
            reverse('riwayat_urls:riwayat_cuti_melahirkan'),
            {'jenis_cuti': 'Cuti Sakit', 'tahun': '2025', 'q': 'Siti'},
        )

        self.assertEqual(list(response.context['cuti_list']), [self.cuti_sakit])

    def test_export_excel_mengikuti_filter_dan_memuat_penempatan(self):
        response = self.client.get(
            reverse('riwayat_urls:data_cuti_export_excel'),
            {'jenis_cuti': 'Cuti Tahunan', 'tahun': '2026'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml.sheet', response['Content-Type'])
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], 'Siti Sehat')
        self.assertEqual(rows[1][4], 'Bidang Pelayanan')
        self.assertEqual(rows[1][5], 'Cuti Tahunan')

    def test_pengguna_biasa_ditolak(self):
        self.client.force_login(self.pegawai)
        response = self.client.get(reverse('riwayat_urls:riwayat_cuti_melahirkan'))
        self.assertEqual(response.status_code, 403)
